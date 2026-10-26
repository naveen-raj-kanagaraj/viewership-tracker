"""Logs for the tracked event: a spreadsheet export (any time) and an automatic summary after each broadcast day."""
import io
import os
import time
from datetime import datetime

import requests

from tracker import events, store

LOG_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "logs")
QUIET_MIN = 45   # a broadcast counts as finished after this long with no readings


def _t(ts):
    return datetime.fromtimestamp(ts, events.REPORT_TZ).strftime("%d %b %Y %I:%M:%S %p") if ts else ""


def export_xlsx(key: str = None) -> bytes:
    """Spreadsheet: Summary, By day, Broadcasts, Streams (each one's peak), Watch parties, Readings (per minute)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    ev = events.get(key)
    st = events.stats(ev["key"])
    wb = Workbook()
    head = Font(bold=True, color="0A0A14")
    fill = PatternFill("solid", fgColor="B6F23A")

    def sheet(title, cols, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(r)
        for i, col in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(48, len(str(col)) + 6))
        return ws

    sp = st.get("peak_split", {})
    sheet("Summary", ["Metric", "Value"], [
        ["Event", ev["name"]], ["Peak concurrent (all streams)", st.get("peak", 0)], ["Peak at (IST)", _t(st.get("peak_at"))],
        ["Official at peak", sp.get("official", 0)], ["Watch parties at peak", sp.get("watch", 0)],
        ["Average concurrent (while official live)", round(st.get("avg_broadcast", 0))],
        ["Official peak", st.get("peak_official", 0)], ["Watch parties peak (combined)", st.get("peak_watch", 0)],
        ["Hours watched", round(st.get("hours_watched", 0))], ["Exported (IST)", _t(time.time())],
        ["Source", f"{', '.join(st.get('platform_names') or ['YouTube'])}, counted from each stream's live counter; recorded while the official broadcast is live"]],
        first=True)
    sheet("By day", ["Day (IST)", "Peak concurrent", "Peak at", "Average concurrent (broadcast)", "Hours watched"],
          [[d["day"], d["peak"], _t(d.get("peak_at")), round(d.get("avg_broadcast", 0)), round(d.get("hw", 0))] for d in st.get("days", [])])
    with store._lock:
        c = events._db()
        ticks = c.execute("SELECT ts, total, official, watch FROM ev_ticks WHERE event=? ORDER BY ts", (ev["key"],)).fetchall()
        streams = c.execute("SELECT vid, channel, kind, title, peak, peak_at, first_seen, last_seen FROM ev_streams "
                            "WHERE event=? AND peak>0 ORDER BY peak DESC", (ev["key"],)).fetchall()
    sess, cur = [], None
    for ts, tot, o, w in ticks:
        if cur and ts - cur[1] <= 3600:
            cur[1] = ts
            if tot > cur[2]:
                cur[2:6] = [tot, ts, o, w]
        else:
            cur = [ts, ts, tot, ts, o, w]
            sess.append(cur)
    sheet("Broadcasts", ["Start (IST)", "End (IST)", "Peak concurrent", "Peak at", "Official at peak", "Watch parties at peak"],
          [[_t(a), _t(b), p, _t(pa), o, w] for a, b, p, pa, o, w in sess])
    sheet("Streams", ["Channel", "Type", "Peak concurrent", "Peak at (IST)", "First seen", "Last seen", "Title", "Link"],
          [[ch, k, p, _t(pa), _t(fs), _t(ls), t, events.yt_link(v)] for v, ch, k, t, p, pa, fs, ls in streams])
    sheet("Watch parties", ["Rank", "Channel", "Peak concurrent", "Peak at (IST)", "Link"],
          [[i, ch, p, _t(pa), events.yt_link(v)] for i, (v, ch, k, t, p, pa, fs, ls) in
           enumerate([s for s in streams if s[2] == "watch party"], 1)])
    per_min = {}
    for ts, tot, o, w in ticks:
        per_min[ts // 60] = (ts, tot, o, w)
    sheet("Readings (per minute)", ["Time (IST)", "Total concurrent", "Official", "Watch parties"],
          [[_t(ts), tot, o, w] for ts, tot, o, w in per_min.values()])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def anomalies_xlsx(key: str, since: float = 0) -> tuple:
    """(xlsx bytes, row count) of sudden viewer jumps/drops per stream."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    with store._lock:
        rows = events._db().execute("SELECT ts, channel, kind, before, after, change, pct, seconds, vid FROM ev_anomalies "
                                    "WHERE event=? AND ts>=? ORDER BY ts", (key, since)).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Anomalies"
    ws.append(["Time (IST)", "Channel", "Type", "Before", "After", "Change", "Change %", "Within (s)", "Link"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="B6F23A")
    for ts, ch, k, b, a, d, p, sec, vid in rows:
        ws.append([_t(ts), ch, k, b, a, d, p, sec, events.yt_link(vid)])
    for col, w in zip("ABCDEFGHI", (24, 26, 12, 10, 10, 10, 10, 10, 32)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(rows)


def check_anomalies(ev: dict):
    """Every 30 min, if new jumps/drops were logged, send the spreadsheet of today's anomalies."""
    key = ev["key"]
    last = store.kv_get(f"ev_anom_sent:{key}", 0)
    if time.time() - last < 1800:
        return
    with store._lock:
        n = events._db().execute("SELECT COUNT(*) FROM ev_anomalies WHERE event=? AND ts>?", (key, last)).fetchone()[0]
    if not n:
        return
    store.kv_set(f"ev_anom_sent:{key}", time.time())
    day0 = datetime.now(events.REPORT_TZ).replace(hour=0, minute=0, second=0, microsecond=0).timestamp()
    xl, total = anomalies_xlsx(key, day0)
    _send(f"{ev['name']}: {n} new sudden viewer jumps/drops ({total} today). Per stream, change within ~15s.",
          xl, f"{key}-anomalies-{datetime.now(events.REPORT_TZ):%Y-%m-%d}.xlsx")


def _send(text: str, doc: bytes = None, name: str = ""):
    token = os.getenv("TELEGRAM_BOT_TOKEN")
    chat = (os.getenv("OWNER_TELEGRAM_ID") or "").split(",")[0].strip()
    if not (token and chat):
        return
    try:
        if doc:
            requests.post(f"https://api.telegram.org/bot{token}/sendDocument", data={"chat_id": chat, "caption": text[:1000]},
                          files={"document": (name, doc)}, timeout=60)
        else:
            requests.post(f"https://api.telegram.org/bot{token}/sendMessage", json={"chat_id": chat, "text": text,
                          "disable_web_page_preview": True}, timeout=20)
    except requests.RequestException:
        pass


def day_summary(ev: dict, lb: dict, st: dict) -> str:
    sp = lb.get("split", {})
    wps = [x for x in lb.get("streams", []) if x["kind"] == "watch party"]
    off = [x for x in lb.get("streams", []) if x["kind"] == "official"]
    tf = lambda ts: datetime.fromtimestamp(ts, events.REPORT_TZ).strftime("%I:%M %p")
    day = datetime.fromtimestamp(lb["start"], events.REPORT_TZ).strftime("%a %d %b %Y")
    lines = [f"{ev['name']} // {day} // broadcast log",
             f"{tf(lb['start'])} to {tf(lb['end'])} IST",
             f"\nPEAK CONCURRENT: {lb['peak']:,} at {tf(lb['peak_at'])} IST",
             f"  official {sp.get('official', 0):,} | watch parties {sp.get('watch', 0):,}"]
    if off:
        lines.append("\nOfficial streams (peak):")
        lines += [f"  {x['peak']:,}  {x['title'][:60]}" for x in off[:4]]
    if wps:
        lines.append("\nWatch parties (peak):")
        lines += [f"  {i}. {x['channel']}  {x['peak']:,} at {tf(x['peak_at'])}" for i, x in enumerate(wps[:6], 1)]
    lines.append(f"\nEvent peak so far: {st.get('peak', 0):,}")
    return "\n".join(lines)


NOT_LIVE_AFTER = int(os.getenv("EVENT_NOT_LIVE_MIN", "10"))   # minutes after the official stream starts


def check_roster(ev: dict):
    """Once per broadcast, NOT_LIVE_AFTER minutes after the official stream starts: tell the owner on Telegram which
    permitted watch-party streamers are still not live, so their link can be fixed or they can be removed."""
    if not events.PRESETS.get(ev.get("key"), {}).get("roster") or events.PRESETS.get(ev.get("key"), {}).get("ended"):
        return
    st = events.stats(ev["key"])
    lb = st.get("last_broadcast") or {}
    if not st.get("official_live") or not lb or time.time() - lb["start"] < NOT_LIVE_AFTER * 60:
        return
    flag = f"ev_notlive:{ev['key']}:{lb['start']}"
    if store.kv_get(flag):
        return
    from tracker import eventpage
    roster = eventpage.data(ev["key"]).get("roster") or []
    missing = [r for r in roster if r.get("status") == "upcoming"]
    store.kv_set(flag, True)
    if not missing:
        _send(f"{ev['name']}: all {len(roster)} watch-party streamers are live.")
        return
    lines = [f"{ev['name']}: {len(missing)} of {len(roster)} watch-party streamers are NOT live "
             f"{NOT_LIVE_AFTER} min after the official stream started:"]
    lines += [f"- {r['name']}: {r.get('url') or ''}" for r in missing]
    lines.append("Send me their live stream link to add it, or tell me to remove them from the watch-party list.")
    _send("\n".join(lines))


def check(ev: dict):
    """Called every minute: once a broadcast has been quiet for QUIET_MIN, save its log and send it to the owner once."""
    st = events.stats(ev["key"])
    lb = st.get("last_broadcast")
    if not lb or st.get("official_live") or time.time() - lb["end"] < QUIET_MIN * 60 or lb["end"] - lb["start"] < 600:
        return
    flag = f"ev_daylog:{ev['key']}:{lb['start']}"
    if store.kv_get(flag):
        return
    store.kv_set(flag, True)
    text = day_summary(ev, lb, st)
    os.makedirs(LOG_DIR, exist_ok=True)
    stamp = datetime.fromtimestamp(lb["start"], events.REPORT_TZ).strftime("%Y-%m-%d_%H%M")
    with open(os.path.join(LOG_DIR, f"{ev['key']}_{stamp}.txt"), "w") as f:
        f.write(text + "\n")
    try:
        xl = export_xlsx(ev["key"])
        with open(os.path.join(LOG_DIR, f"{ev['key']}_{stamp}.xlsx"), "wb") as f:
            f.write(xl)
    except Exception as e:
        xl = None
        print(f"Event export failed: {str(e)[:150]}")
    _send(text)
    if xl:
        _send(f"{ev['name']} full data so far (spreadsheet)", xl, f"{ev['key']}_{stamp}.xlsx")
